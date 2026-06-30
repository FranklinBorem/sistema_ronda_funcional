"""
services/relatorio_conferencia.py — Grupo Ronda · Gerador de Relatório Técnico

Gera relatório de monitoramento de NVRs e câmeras em PDF ou Excel.
Baseado nos modelos: NvrStatusLog, CameraStatusLog, MonitorEvent.

Uso:
    from services.relatorio_conferencia import gerar_relatorio
    pdf_bytes = gerar_relatorio(dt_inicio, dt_fim, fmt="pdf")
    xlsx_bytes = gerar_relatorio(dt_inicio, dt_fim, fmt="xlsx")
"""

from __future__ import annotations

import io
from datetime import datetime, timedelta
from typing import Literal

# ── ReportLab ────────────────────────────────────────────────────────────────
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    HRFlowable, PageBreak, Paragraph, SimpleDocTemplate,
    Spacer, Table, TableStyle,
)

# ── OpenPyXL ─────────────────────────────────────────────────────────────────
import openpyxl
from openpyxl.styles import (
    Alignment, Border, Font, PatternFill, Side,
)
from openpyxl.utils import get_column_letter


# ──────────────────────────────────────────────────────────────────────────────
# Cores Grupo Ronda
# ──────────────────────────────────────────────────────────────────────────────
GOLD   = colors.HexColor("#F5C400")
BLACK  = colors.HexColor("#0A0A0A")
DARK   = colors.HexColor("#1A1A1A")
GRAY   = colors.HexColor("#4A4A4A")
LGRAY  = colors.HexColor("#E8E8E8")
GREEN  = colors.HexColor("#22C55E")
RED    = colors.HexColor("#EF4444")
ORANGE = colors.HexColor("#F59E0B")
WHITE  = colors.white

GOLD_HEX   = "F5C400"
BLACK_HEX  = "0A0A0A"
RED_HEX    = "EF4444"
GREEN_HEX  = "22C55E"
ORANGE_HEX = "F59E0B"
LGRAY_HEX  = "F3F3F3"
DGRAY_HEX  = "4A4A4A"


# ──────────────────────────────────────────────────────────────────────────────
# Estilos ReportLab
# ──────────────────────────────────────────────────────────────────────────────

def _estilos():
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "rg_title", fontSize=16, fontName="Helvetica-Bold",
            textColor=BLACK, alignment=TA_CENTER, spaceAfter=2,
        ),
        "subtitle": ParagraphStyle(
            "rg_subtitle", fontSize=10, fontName="Helvetica",
            textColor=GRAY, alignment=TA_CENTER, spaceAfter=2,
        ),
        "section": ParagraphStyle(
            "rg_section", fontSize=9, fontName="Helvetica-Bold",
            textColor=BLACK, spaceBefore=8, spaceAfter=4,
        ),
        "body": ParagraphStyle(
            "rg_body", fontSize=8, fontName="Helvetica",
            textColor=GRAY, spaceAfter=3, leading=12,
        ),
        "label": ParagraphStyle(
            "rg_label", fontSize=7, fontName="Helvetica-Bold",
            textColor=GRAY, alignment=TA_CENTER,
        ),
        "value": ParagraphStyle(
            "rg_value", fontSize=9, fontName="Helvetica-Bold",
            textColor=BLACK, alignment=TA_CENTER,
        ),
        "footer": ParagraphStyle(
            "rg_footer", fontSize=7, fontName="Helvetica",
            textColor=GRAY, alignment=TA_CENTER,
        ),
    }


# ──────────────────────────────────────────────────────────────────────────────
# Consultas ao banco
# ──────────────────────────────────────────────────────────────────────────────

def _buscar_dados(dt_inicio: datetime, dt_fim: datetime) -> dict:
    from models.monitoring import CameraStatusLog, MonitorEvent
    from models.nvr_status_log import NvrStatusLog

    # Logs de NVR no período
    nvr_logs = (
        NvrStatusLog.query
        .filter(NvrStatusLog.timestamp >= dt_inicio,
                NvrStatusLog.timestamp <= dt_fim)
        .order_by(NvrStatusLog.nvr_id, NvrStatusLog.timestamp)
        .all()
    )

    # Logs de câmera no período
    cam_logs = (
        CameraStatusLog.query
        .filter(CameraStatusLog.timestamp >= dt_inicio,
                CameraStatusLog.timestamp <= dt_fim)
        .order_by(CameraStatusLog.nvr_id,
                  CameraStatusLog.channel_id,
                  CameraStatusLog.timestamp)
        .all()
    )

    # Eventos que se sobrepõem ao período
    eventos = (
        MonitorEvent.query
        .filter(
            MonitorEvent.aberto_em <= dt_fim,
            (MonitorEvent.resolvido_em == None) |
            (MonitorEvent.resolvido_em >= dt_inicio),
        )
        .order_by(MonitorEvent.aberto_em)
        .all()
    )

    return {"nvr_logs": nvr_logs, "cam_logs": cam_logs, "eventos": eventos}


def _calcular_resumo(dados: dict, dt_inicio: datetime, dt_fim: datetime) -> dict:
    """Calcula métricas agregadas para o relatório."""
    nvr_logs = dados["nvr_logs"]
    cam_logs = dados["cam_logs"]
    eventos  = dados["eventos"]

    periodo_h = max((dt_fim - dt_inicio).total_seconds() / 3600, 0.01)

    # NVRs únicos
    nvrs = {}
    for log in nvr_logs:
        if log.nvr_id not in nvrs:
            nvrs[log.nvr_id] = {"nome": log.nvr_nome, "logs": []}
        nvrs[log.nvr_id]["logs"].append(log)

    # Uptime por NVR
    for nvr_id, info in nvrs.items():
        logs = info["logs"]
        total = len(logs)
        ok = sum(1 for l in logs if l.health == "OK")
        info["uptime_pct"] = round(ok / total * 100, 1) if total else 0
        info["total_polls"] = total
        info["polls_ok"] = ok

    # Câmeras únicas
    cameras = {}
    for log in cam_logs:
        key = (log.nvr_id, log.channel_id)
        if key not in cameras:
            cameras[key] = {"nome": log.channel_name, "nvr_id": log.nvr_id, "logs": []}
        cameras[key]["logs"].append(log)

    for key, info in cameras.items():
        logs = info["logs"]
        total = len(logs)
        online = sum(1 for l in logs if l.online)
        info["uptime_pct"] = round(online / total * 100, 1) if total else 0
        info["total_polls"] = total
        info["polls_online"] = online

    # Eventos por categoria
    nvr_eventos  = [e for e in eventos if e.alvo_tipo == "nvr"]
    cam_eventos  = [e for e in eventos if e.alvo_tipo == "camera"]
    ativos       = [e for e in eventos if e.resolvido_em is None]

    return {
        "nvrs": nvrs,
        "cameras": cameras,
        "nvr_eventos": nvr_eventos,
        "cam_eventos": cam_eventos,
        "ativos": ativos,
        "periodo_h": periodo_h,
        "total_eventos": len(eventos),
        "total_nvrs": len(nvrs),
        "total_cameras": len(cameras),
        "cameras_problema": sum(1 for info in cameras.values() if info["uptime_pct"] < 100),
    }


def _fmt_dur(minutos: float | None) -> str:
    if minutos is None:
        return "–"
    if minutos < 60:
        return f"{int(minutos)}min"
    h = int(minutos // 60)
    m = int(minutos % 60)
    return f"{h}h {m}min"


def _fmt_dt(dt: datetime | None) -> str:
    if dt is None:
        return "Em aberto"
    return dt.strftime("%d/%m/%Y %H:%M")


# ──────────────────────────────────────────────────────────────────────────────
# Gerador PDF
# ──────────────────────────────────────────────────────────────────────────────

def _gerar_pdf(dt_inicio: datetime, dt_fim: datetime, dados: dict, resumo: dict) -> bytes:
    buf = io.BytesIO()
    W, H = A4
    margin = 18 * mm

    estilos = _estilos()
    story = []

    # ── Cabeçalho (papel timbrado) ────────────────────────────────────────
    # Faixa dourada superior
    header_table = Table(
        [[
            Paragraph("<b>GRUPO<br/>RONDA</b>",
                      ParagraphStyle("logo", fontSize=14, fontName="Helvetica-Bold",
                                     textColor=BLACK, leading=16)),
            Paragraph(
                "RELATÓRIO TÉCNICO OPERACIONAL<br/>"
                "<font size=9>Monitoramento de NVRs e Câmeras</font>",
                ParagraphStyle("htitle", fontSize=12, fontName="Helvetica-Bold",
                               textColor=WHITE, alignment=TA_CENTER, leading=16)
            ),
            Paragraph(
                f"<font size=7>Emitido em<br/>{datetime.now().strftime('%d/%m/%Y %H:%M')}</font>",
                ParagraphStyle("hdate", fontSize=7, fontName="Helvetica",
                               textColor=WHITE, alignment=TA_RIGHT)
            ),
        ]],
        colWidths=[40*mm, W - 2*margin - 80*mm, 40*mm],
    )
    header_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, 0), GOLD),
        ("BACKGROUND", (1, 0), (2, 0), BLACK),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING",  (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING",   (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING",(0, 0), (-1, -1), 10),
    ]))
    story.append(header_table)
    story.append(Spacer(1, 4*mm))

    # ── Período ───────────────────────────────────────────────────────────
    story.append(Paragraph(
        f"Período: {dt_inicio.strftime('%d/%m/%Y %H:%M')} — {dt_fim.strftime('%d/%m/%Y %H:%M')}",
        ParagraphStyle("periodo", fontSize=9, fontName="Helvetica-Bold",
                       textColor=GRAY, alignment=TA_CENTER),
    ))
    story.append(Spacer(1, 3*mm))
    story.append(HRFlowable(width="100%", thickness=1, color=GOLD))
    story.append(Spacer(1, 4*mm))

    # ── Resumo executivo ──────────────────────────────────────────────────
    story.append(Paragraph("1. RESUMO EXECUTIVO", estilos["section"]))

    # Cards de métricas
    cards = [
        [
            Paragraph(str(resumo["total_nvrs"]), estilos["value"]),
            Paragraph(str(resumo["total_cameras"]), estilos["value"]),
            Paragraph(str(resumo["total_eventos"]), estilos["value"]),
            Paragraph(str(len(resumo["ativos"])), estilos["value"]),
            Paragraph(str(resumo["cameras_problema"]), estilos["value"]),
        ],
        [
            Paragraph("NVRs Monitorados", estilos["label"]),
            Paragraph("Câmeras Monitoradas", estilos["label"]),
            Paragraph("Total de Eventos", estilos["label"]),
            Paragraph("Eventos Ativos", estilos["label"]),
            Paragraph("Câmeras c/ Falha", estilos["label"]),
        ],
    ]
    cw = (W - 2*margin) / 5
    cards_table = Table(cards, colWidths=[cw]*5)
    cards_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), LGRAY),
        ("BACKGROUND", (3, 0), (3, 0), RED if resumo["ativos"] else GREEN),
        ("BACKGROUND", (4, 0), (4, 0), ORANGE if resumo["cameras_problema"] else GREEN),
        ("TEXTCOLOR",  (3, 0), (3, 0), WHITE),
        ("TEXTCOLOR",  (4, 0), (4, 0), WHITE),
        ("BOX", (0, 0), (-1, -1), 0.5, LGRAY),
        ("INNERGRID", (0, 0), (-1, -1), 0.3, LGRAY),
        ("TOPPADDING",   (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING",(0, 0), (-1, -1), 6),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(cards_table)
    story.append(Spacer(1, 4*mm))

    # ── Status por NVR ────────────────────────────────────────────────────
    story.append(Paragraph("2. STATUS DOS NVRs", estilos["section"]))

    nvr_rows = [[
        Paragraph("NVR", estilos["label"]),
        Paragraph("Total Polls", estilos["label"]),
        Paragraph("Polls OK", estilos["label"]),
        Paragraph("Uptime %", estilos["label"]),
        Paragraph("Câmeras Online (último)", estilos["label"]),
    ]]
    for nvr_id, info in resumo["nvrs"].items():
        logs = info["logs"]
        ultimo = logs[-1] if logs else None
        uptime = info["uptime_pct"]
        cor = GREEN if uptime >= 99 else (ORANGE if uptime >= 90 else RED)
        nvr_rows.append([
            Paragraph(info["nome"], estilos["body"]),
            Paragraph(str(info["total_polls"]), estilos["body"]),
            Paragraph(str(info["polls_ok"]), estilos["body"]),
            Paragraph(
                f'<font color="#{("22C55E" if uptime>=99 else ("F59E0B" if uptime>=90 else "EF4444"))}">'
                f'<b>{uptime}%</b></font>',
                ParagraphStyle("up", fontSize=8, fontName="Helvetica-Bold", alignment=TA_CENTER)
            ),
            Paragraph(
                f'{ultimo.cameras_online}/{ultimo.cameras_total}' if ultimo else '–',
                estilos["body"]
            ),
        ])

    cw_nvr = [(W - 2*margin) * x for x in [0.35, 0.13, 0.13, 0.13, 0.26]]
    nvr_table = Table(nvr_rows, colWidths=cw_nvr)
    nvr_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), BLACK),
        ("TEXTCOLOR",  (0, 0), (-1, 0), WHITE),
        ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE",   (0, 0), (-1, 0), 7),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [WHITE, LGRAY]),
        ("BOX", (0, 0), (-1, -1), 0.5, GRAY),
        ("INNERGRID", (0, 0), (-1, -1), 0.3, LGRAY),
        ("TOPPADDING",   (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING",(0, 0), (-1, -1), 5),
        ("LEFTPADDING",  (0, 0), (-1, -1), 5),
        ("ALIGN", (1, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(nvr_table)
    story.append(Spacer(1, 4*mm))

    # ── Eventos de falha ──────────────────────────────────────────────────
    story.append(Paragraph("3. EVENTOS DE FALHA NO PERÍODO", estilos["section"]))

    if not dados["eventos"]:
        story.append(Paragraph(
            "Nenhum evento de falha registrado no período.",
            ParagraphStyle("ok", fontSize=8, fontName="Helvetica",
                           textColor=GREEN, spaceAfter=6),
        ))
    else:
        ev_rows = [[
            Paragraph("Dispositivo", estilos["label"]),
            Paragraph("Tipo", estilos["label"]),
            Paragraph("Severidade", estilos["label"]),
            Paragraph("Início", estilos["label"]),
            Paragraph("Fim", estilos["label"]),
            Paragraph("Duração", estilos["label"]),
            Paragraph("Status", estilos["label"]),
        ]]
        for e in dados["eventos"]:
            ativo = e.resolvido_em is None
            sev_cor = {"critica": "EF4444", "alta": "F59E0B",
                       "media": "3B82F6", "baixa": "22C55E"}.get(e.severidade, "888888")
            ev_rows.append([
                Paragraph(e.nome_exibicao[:35], estilos["body"]),
                Paragraph(e.alvo_tipo.upper(), estilos["body"]),
                Paragraph(
                    f'<font color="#{sev_cor}"><b>{e.severidade.upper()}</b></font>',
                    ParagraphStyle("sev", fontSize=7, fontName="Helvetica-Bold", alignment=TA_CENTER)
                ),
                Paragraph(_fmt_dt(e.aberto_em), estilos["body"]),
                Paragraph(_fmt_dt(e.resolvido_em), estilos["body"]),
                Paragraph(_fmt_dur(e.duracao_minutos), estilos["body"]),
                Paragraph(
                    '<font color="#EF4444"><b>ATIVO</b></font>' if ativo
                    else '<font color="#22C55E"><b>RESOLVIDO</b></font>',
                    ParagraphStyle("st", fontSize=7, fontName="Helvetica-Bold", alignment=TA_CENTER)
                ),
            ])

        cw_ev = [(W - 2*margin) * x for x in [0.22, 0.08, 0.10, 0.16, 0.16, 0.12, 0.16]]
        ev_table = Table(ev_rows, colWidths=cw_ev)
        ev_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), BLACK),
            ("TEXTCOLOR",  (0, 0), (-1, 0), WHITE),
            ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE",   (0, 0), (-1, 0), 7),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [WHITE, LGRAY]),
            ("BOX", (0, 0), (-1, -1), 0.5, GRAY),
            ("INNERGRID", (0, 0), (-1, -1), 0.3, LGRAY),
            ("TOPPADDING",   (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING",(0, 0), (-1, -1), 4),
            ("LEFTPADDING",  (0, 0), (-1, -1), 4),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        story.append(ev_table)

    story.append(Spacer(1, 4*mm))

    # ── Câmeras com falha ─────────────────────────────────────────────────
    cams_problema = {k: v for k, v in resumo["cameras"].items() if v["uptime_pct"] < 100}
    if cams_problema:
        story.append(Paragraph("4. CÂMERAS COM FALHA NO PERÍODO", estilos["section"]))

        cam_rows = [[
            Paragraph("NVR", estilos["label"]),
            Paragraph("Canal", estilos["label"]),
            Paragraph("Nome", estilos["label"]),
            Paragraph("Polls", estilos["label"]),
            Paragraph("Online", estilos["label"]),
            Paragraph("Uptime %", estilos["label"]),
        ]]
        for (nvr_id, ch_id), info in sorted(cams_problema.items()):
            uptime = info["uptime_pct"]
            cam_rows.append([
                Paragraph(info["nvr_id"], estilos["body"]),
                Paragraph(str(ch_id), estilos["body"]),
                Paragraph(info["nome"][:30], estilos["body"]),
                Paragraph(str(info["total_polls"]), estilos["body"]),
                Paragraph(str(info["polls_online"]), estilos["body"]),
                Paragraph(
                    f'<font color="#{"EF4444" if uptime<90 else "F59E0B"}"><b>{uptime}%</b></font>',
                    ParagraphStyle("cu", fontSize=8, fontName="Helvetica-Bold", alignment=TA_CENTER)
                ),
            ])

        cw_cam = [(W - 2*margin) * x for x in [0.20, 0.08, 0.30, 0.10, 0.10, 0.22]]
        cam_table = Table(cam_rows, colWidths=cw_cam)
        cam_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), BLACK),
            ("TEXTCOLOR",  (0, 0), (-1, 0), WHITE),
            ("FONTNAME",   (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE",   (0, 0), (-1, 0), 7),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [WHITE, LGRAY]),
            ("BOX", (0, 0), (-1, -1), 0.5, GRAY),
            ("INNERGRID", (0, 0), (-1, -1), 0.3, LGRAY),
            ("TOPPADDING",   (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING",(0, 0), (-1, -1), 4),
            ("LEFTPADDING",  (0, 0), (-1, -1), 4),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        story.append(cam_table)

    # ── Rodapé ────────────────────────────────────────────────────────────
    story.append(Spacer(1, 6*mm))
    story.append(HRFlowable(width="100%", thickness=0.5, color=GOLD))
    story.append(Spacer(1, 2*mm))
    story.append(Paragraph(
        "Grupo Ronda Porteiros LTDA · Central de Monitoramento Ronda · "
        f"Documento gerado em {datetime.now().strftime('%d/%m/%Y às %H:%M')}",
        estilos["footer"],
    ))

    def _add_page_num(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(GRAY)
        canvas.drawRightString(W - margin, 10*mm, f"Página {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=margin, rightMargin=margin,
        topMargin=10*mm, bottomMargin=18*mm,
    )
    doc.build(story, onFirstPage=_add_page_num, onLaterPages=_add_page_num)
    return buf.getvalue()


# ──────────────────────────────────────────────────────────────────────────────
# Gerador Excel
# ──────────────────────────────────────────────────────────────────────────────

def _hdr_style(ws, row, cols, val, fill_hex=BLACK_HEX):
    fill = PatternFill("solid", fgColor=fill_hex)
    font = Font(bold=True, color="FFFFFF", size=9)
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.cell(row=row, column=1).value = val
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)


def _row_style(cell, bold=False, color=None, align="left", bg=None):
    cell.font = Font(bold=bold, size=8, color=color or "000000")
    cell.alignment = Alignment(horizontal=align, vertical="center")
    if bg:
        cell.fill = PatternFill("solid", fgColor=bg)


def _gerar_xlsx(dt_inicio: datetime, dt_fim: datetime, dados: dict, resumo: dict) -> bytes:
    wb = openpyxl.Workbook()

    thin = Side(style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ── Aba 1: Resumo ──────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Resumo"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 20

    gold_fill = PatternFill("solid", fgColor=GOLD_HEX)
    black_fill = PatternFill("solid", fgColor=BLACK_HEX)

    # Cabeçalho
    ws.merge_cells("A1:B1")
    ws["A1"] = "GRUPO RONDA — RELATÓRIO DE MONITORAMENTO"
    ws["A1"].font = Font(bold=True, size=13, color="FFFFFF")
    ws["A1"].fill = black_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:B2")
    ws["A2"] = f"Período: {dt_inicio.strftime('%d/%m/%Y %H:%M')} — {dt_fim.strftime('%d/%m/%Y %H:%M')}"
    ws["A2"].font = Font(bold=True, size=9, color=BLACK_HEX)
    ws["A2"].fill = PatternFill("solid", fgColor=GOLD_HEX)
    ws["A2"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 16

    # Métricas
    metricas = [
        ("NVRs Monitorados", resumo["total_nvrs"]),
        ("Câmeras Monitoradas", resumo["total_cameras"]),
        ("Total de Eventos", resumo["total_eventos"]),
        ("Eventos Ativos", len(resumo["ativos"])),
        ("Câmeras com Falha", resumo["cameras_problema"]),
        ("Período (horas)", round(resumo["periodo_h"], 1)),
    ]
    for i, (label, val) in enumerate(metricas, start=4):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True, size=9)
        c = ws.cell(row=i, column=2, value=val)
        c.font = Font(bold=True, size=11)
        c.alignment = Alignment(horizontal="center")
        if label == "Eventos Ativos" and val > 0:
            c.font = Font(bold=True, size=11, color=RED_HEX)
        elif label == "Câmeras com Falha" and val > 0:
            c.font = Font(bold=True, size=11, color=ORANGE_HEX)

    # ── Aba 2: NVRs ───────────────────────────────────────────────────────
    ws2 = wb.create_sheet("NVRs")
    hdrs = ["NVR ID", "Nome", "Total Polls", "Polls OK", "Uptime %",
            "Câm Online (último)", "Câm Offline (último)"]
    for i, h in enumerate(hdrs, 1):
        c = ws2.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=9)
        c.fill = black_fill
        c.alignment = Alignment(horizontal="center")
        ws2.column_dimensions[get_column_letter(i)].width = [15,25,13,10,10,18,18][i-1]

    for row_i, (nvr_id, info) in enumerate(resumo["nvrs"].items(), start=2):
        logs = info["logs"]
        ultimo = logs[-1] if logs else None
        uptime = info["uptime_pct"]
        vals = [
            nvr_id, info["nome"], info["total_polls"], info["polls_ok"],
            uptime,
            ultimo.cameras_online if ultimo else 0,
            ultimo.cameras_offline if ultimo else 0,
        ]
        bg = LGRAY_HEX if row_i % 2 == 0 else "FFFFFF"
        for col_i, val in enumerate(vals, 1):
            c = ws2.cell(row=row_i, column=col_i, value=val)
            c.fill = PatternFill("solid", fgColor=bg)
            c.alignment = Alignment(horizontal="center" if col_i > 1 else "left", vertical="center")
            c.border = border
            if col_i == 5:  # uptime
                if uptime >= 99:
                    c.font = Font(bold=True, color=GREEN_HEX)
                elif uptime >= 90:
                    c.font = Font(bold=True, color=ORANGE_HEX)
                else:
                    c.font = Font(bold=True, color=RED_HEX)

    # ── Aba 3: Eventos ────────────────────────────────────────────────────
    ws3 = wb.create_sheet("Eventos")
    hdrs3 = ["Dispositivo", "NVR ID", "Canal", "Tipo", "Severidade",
             "Início", "Fim", "Duração", "Status"]
    for i, h in enumerate(hdrs3, 1):
        c = ws3.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=9)
        c.fill = black_fill
        c.alignment = Alignment(horizontal="center")
        ws3.column_dimensions[get_column_letter(i)].width = [28,15,8,10,12,18,18,12,12][i-1]

    for row_i, e in enumerate(dados["eventos"], start=2):
        ativo = e.resolvido_em is None
        vals = [
            e.nome_exibicao, e.nvr_id, e.channel_id or "–",
            e.alvo_tipo.upper(), e.severidade.upper(),
            e.aberto_em.strftime("%d/%m/%Y %H:%M"),
            e.resolvido_em.strftime("%d/%m/%Y %H:%M") if e.resolvido_em else "Em aberto",
            _fmt_dur(e.duracao_minutos),
            "ATIVO" if ativo else "RESOLVIDO",
        ]
        bg = LGRAY_HEX if row_i % 2 == 0 else "FFFFFF"
        for col_i, val in enumerate(vals, 1):
            c = ws3.cell(row=row_i, column=col_i, value=val)
            c.fill = PatternFill("solid", fgColor=bg)
            c.alignment = Alignment(horizontal="center" if col_i > 1 else "left", vertical="center")
            c.border = border
            if col_i == 9:  # status
                c.font = Font(bold=True,
                              color=RED_HEX if ativo else GREEN_HEX)
            if col_i == 5:  # severidade
                sev_col = {"CRITICA": RED_HEX, "ALTA": ORANGE_HEX,
                           "MEDIA": "3B82F6", "BAIXA": GREEN_HEX}.get(str(val), DGRAY_HEX)
                c.font = Font(bold=True, color=sev_col)

    # ── Aba 4: Câmeras ────────────────────────────────────────────────────
    ws4 = wb.create_sheet("Câmeras")
    hdrs4 = ["NVR ID", "Canal", "Nome", "Total Polls", "Online", "Uptime %"]
    for i, h in enumerate(hdrs4, 1):
        c = ws4.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=9)
        c.fill = black_fill
        c.alignment = Alignment(horizontal="center")
        ws4.column_dimensions[get_column_letter(i)].width = [18,8,30,13,10,10][i-1]

    for row_i, ((nvr_id, ch_id), info) in enumerate(
        sorted(resumo["cameras"].items()), start=2
    ):
        uptime = info["uptime_pct"]
        vals = [nvr_id, ch_id, info["nome"],
                info["total_polls"], info["polls_online"], uptime]
        bg = LGRAY_HEX if row_i % 2 == 0 else "FFFFFF"
        for col_i, val in enumerate(vals, 1):
            c = ws4.cell(row=row_i, column=col_i, value=val)
            c.fill = PatternFill("solid", fgColor=bg)
            c.alignment = Alignment(horizontal="center" if col_i != 3 else "left",
                                    vertical="center")
            c.border = border
            if col_i == 6:
                c.font = Font(bold=True,
                              color=GREEN_HEX if uptime >= 99
                              else (ORANGE_HEX if uptime >= 90 else RED_HEX))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ──────────────────────────────────────────────────────────────────────────────
# Ponto de entrada público
# ──────────────────────────────────────────────────────────────────────────────

def gerar_relatorio(
    dt_inicio: datetime,
    dt_fim: datetime,
    fmt: Literal["pdf", "xlsx"] = "pdf",
) -> bytes:
    """
    Gera relatório de monitoramento no formato especificado.

    Args:
        dt_inicio: início do período (datetime, fuso local)
        dt_fim:    fim do período (datetime, fuso local)
        fmt:       "pdf" ou "xlsx"

    Returns:
        bytes do arquivo gerado
    """
    dados  = _buscar_dados(dt_inicio, dt_fim)
    resumo = _calcular_resumo(dados, dt_inicio, dt_fim)

    if fmt == "xlsx":
        return _gerar_xlsx(dt_inicio, dt_fim, dados, resumo)
    return _gerar_pdf(dt_inicio, dt_fim, dados, resumo)
